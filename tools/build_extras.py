#!/usr/bin/env python
"""Highlights, revision cover letter and the numbered submission package.
Usage: python -I tools/build_extras.py <out_dir> <package_dir>"""
import json
import os
import shutil
import sys

sys.path.insert(0, "tools")
import docbuild as DB  # noqa: E402
from docx.shared import Pt  # noqa: E402

COVER = [
    "Dear Editor,",
    "We are pleased to submit the revised version of our manuscript “Pyrin versus NLRP3: Discriminative Cell-State Mapping of Inflammasome Sensor Programs in Human Atherosclerotic Plaque” for consideration in Cytokine. We thank you and the three reviewers for their constructive comments.",
    "To address the reviewers' requests, we rebuilt the single-cell analysis from the raw GEO deposits in a fully reproducible pipeline, added doublet removal, ambient-RNA correction and explicit sequencing-depth models, sub-clustered the myeloid compartment, split the shared-effector module into cytokine and lysis arms, restructured the dynamic model, regenerated all figures from code at 600 dpi and added the requested references. During this re-analysis we identified and corrected an error in the assignment of core and adjacent sample labels in the original analysis; the corrected core–adjacent results and their interpretation are described transparently in the manuscript and the response letter. The principal conclusions are now that MEFV expression in human carotid plaque is sparse and concentrated in classical monocytes, that this pattern is robust to the technical controls, and that bulk instability signals largely reflect myeloid composition. The study remains explicitly hypothesis-generating.",
    "A point-by-point response, a version of the manuscript with tracked changes, a clean version, revised Highlights, the revised Supplementary Material and Supplementary Tables, and high-resolution figures are provided. All analysis code is available in the public repository cited in the manuscript.",
    "All authors have approved the revised manuscript. The work is not under consideration elsewhere, and the authors declare no conflicts of interest.",
    "Sincerely,",
    "Remziye Dogan, MD, on behalf of all authors",
]
PACKAGE = [("Cover_letter_revision.docx", "00_Cover_letter_revision.docx"),
           ("Response_to_Reviewers.docx", "01_Response_to_Reviewers.docx"),
           ("Manuscript_revised_tracked_changes.docx", "02_Manuscript_revised_tracked_changes.docx"),
           ("Manuscript_revised_clean.docx", "03_Manuscript_revised_clean.docx"),
           ("Highlights.docx", "04_Highlights.docx"),
           ("Supplementary_Material_revised.docx", "05_Supplementary_Material.docx"),
           ("Supplementary_Tables_revised.xlsx", "06_Supplementary_Tables.xlsx")]
FIGS = ["Figure1_schematic", "Figure2_compartments_depth", "Figure3_myeloid_subsets", "Figure4_regional_bulk",
        "Figure5_two_signal_model", "Figure6_priority"]


def main():
    out, pkg = sys.argv[1:3]
    hl = json.load(open("../manuscript/highlights.json", encoding="utf-8"))
    d = DB.new_document(line_numbers=False)
    d.styles["Normal"].paragraph_format.line_spacing = 1.15
    d.add_paragraph("Highlights", style="Heading 1")
    for h in hl:
        p = d.add_paragraph(style="List Bullet")
        for t, f in DB.parse_inline(h.replace("MEFV", "*MEFV*")):
            DB._fmt_run(p.add_run(t), f)
    d.save(os.path.join(out, "Highlights.docx"))
    d = DB.new_document(line_numbers=False)
    d.styles["Normal"].paragraph_format.line_spacing = 1.15
    d.styles["Normal"].paragraph_format.space_after = Pt(8)
    for c in COVER:
        d.add_paragraph(c)
    d.save(os.path.join(out, "Cover_letter_revision.docx"))
    os.makedirs(os.path.join(pkg, "Figures"), exist_ok=True)
    for a, b in PACKAGE:
        shutil.copy2(os.path.join(out, a), os.path.join(pkg, b))
    F = "results/revision/figures"
    for i, name in enumerate(FIGS, 1):
        for ext in ("tiff", "pdf"):
            shutil.copy2(f"{F}/{name}.{ext}", f"{pkg}/Figures/Figure{i}.{ext}")
    for ext in ("tiff", "pdf", "png"):
        shutil.copy2(f"{F}/Graphical_abstract.{ext}", f"{pkg}/Figures/Graphical_abstract.{ext}")
    print("package refreshed")


if __name__ == "__main__":
    main()
