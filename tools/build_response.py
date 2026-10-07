#!/usr/bin/env python
"""Render the point-by-point response letter from response_template.md, filling numbers from
values_used.json (written by build_manuscript.py) and citing references with the manuscript's
numbering. Usage: python -I tools/build_response.py <response_template.md> <manuscript_template.md> <out_dir>
"""
import json
import os
import re
import sys

sys.path.insert(0, "tools")
import docbuild as DB  # noqa: E402
from docx.shared import Pt, RGBColor, Inches  # noqa: E402


def manuscript_order(ms_text):
    order = []
    for grp in re.findall(r"\[(@[^\]]+)\]", ms_text):
        for k in re.findall(r"@(\w+)", grp):
            if k not in order:
                order.append(k)
    return {k: i + 1 for i, k in enumerate(order)}


def main():
    tpl, ms_tpl, outdir = sys.argv[1:4]
    V = json.load(open(os.path.join(outdir, "values_used.json"), encoding="utf-8"))
    src = open(tpl, encoding="utf-8").read()
    miss = sorted(set(re.findall(r"\{\{(\w+)\}\}", src)) - set(V))
    assert not miss, miss
    for k, v in V.items():
        src = src.replace("{{" + k + "}}", v)
    num = manuscript_order(open(ms_tpl, encoding="utf-8").read())

    def rep(m):
        ns = sorted(set(num[k] for k in re.findall(r"@(\w+)", m.group(1))))
        out, i = [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            out.append(f"{ns[i]}\u2013{ns[j]}" if j - i >= 2 else ",".join(str(x) for x in ns[i:j + 1]))
            i = j + 1
        return "[" + ",".join(out) + "]"
    src = re.sub(r"\[(@[^\]]+)\]", rep, src)
    open(os.path.join(outdir, "Response_to_Reviewers.md"), "w", encoding="utf-8").write(src)
    d = DB.new_document(line_numbers=False)
    st = d.styles["Normal"]
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(6)
    for raw in re.split(r"\n\s*\n", src.strip()):
        lines = [l for l in raw.strip().split("\n") if l.strip()]
        if all(l.startswith("- ") for l in lines):
            for l in lines:
                p = d.add_paragraph(style="List Bullet")
                for t, f in DB.parse_inline(l[2:]):
                    DB._fmt_run(p.add_run(t), f)
            continue
        if all(re.match(r"^\d+\. ", l) for l in lines):
            for l in lines:
                p = d.add_paragraph(style="List Number")
                for t, f in DB.parse_inline(re.sub(r"^\d+\. ", "", l)):
                    DB._fmt_run(p.add_run(t), f)
            continue
        text = " ".join(lines)
        if text.startswith("> "):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            for t, f in DB.parse_inline(text[2:]):
                run = p.add_run(t)
                DB._fmt_run(run, f)
                run.italic = True
                run.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", text)
        if m:
            p = d.add_paragraph(style=f"Heading {len(m.group(1))}")
            for t, f in DB.parse_inline(m.group(2)):
                DB._fmt_run(p.add_run(t), f)
            continue
        p = d.add_paragraph()
        for t, f in DB.parse_inline(text):
            DB._fmt_run(p.add_run(t), f)
    d.save(os.path.join(outdir, "Response_to_Reviewers.docx"))
    print("response words:", len(re.sub(r"[#>*]", "", src).split()))


if __name__ == "__main__":
    main()
