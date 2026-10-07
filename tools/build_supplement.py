#!/usr/bin/env python
"""Build the revised Supplementary Tables (.xlsx, one worksheet per table) and the revised
Supplementary Material (.docx: methods incl. model equations, results, figures S1-S8 with legends,
table index and candidate summaries). Usage: python -I tools/build_supplement.py <out_dir>
"""
import os
import sys

import pandas as pd
import yaml

sys.path.insert(0, "tools")
sys.path.insert(0, "src")
import docbuild as DB  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

T = "results/revision/tables"
FIG = "results/revision/figures/supplementary"
SIG = {
    "Mono_classical": ["CD14", "FCN1", "VCAN", "S100A8", "S100A9", "S100A12", "SELL"],
    "Mono_nonclassical": ["FCGR3A", "CX3CR1", "CDKN1C", "LST1", "MS4A7"],
    "Mono_inflammatory_IL1B": ["IL1B", "CXCL8", "TNF", "CCL3", "CCL4", "NLRP3", "EREG", "PTGS2"],
    "Mac_C1Q": ["C1QA", "C1QB", "C1QC", "APOE", "CD163"],
    "Mac_TREM2_lipid": ["TREM2", "SPP1", "GPNMB", "LPL", "CD9", "FABP5", "APOC1", "ACP5"],
    "Mac_LYVE1_resident": ["LYVE1", "F13A1", "FOLR2", "SELENOP", "MRC1", "STAB1"],
    "cDC1": ["CLEC9A", "XCR1", "CADM1", "IDO1"], "cDC2": ["CD1C", "FCER1A", "CLEC10A", "CD1E"],
    "mregDC_LAMP3": ["LAMP3", "CCR7", "FSCN1", "CCL19", "CCL22"], "pDC": ["LILRA4", "CLEC4C", "IL3RA", "JCHAIN", "TCF4"],
    "Neutrophil": ["CSF3R", "FCGR3B", "CXCR2", "CMTM2", "G0S2", "S100P", "PROK2"], "Proliferating": ["MKI67", "TOP2A", "STMN1"],
    "T_contaminant": ["CD3E", "CD3D", "TRAC"], "SMC_contaminant": ["ACTA2", "TAGLN", "MYH11"], "EC_contaminant": ["PECAM1", "VWF", "CDH5"],
}
BROAD = {"Endothelial": ["PECAM1", "VWF", "CDH5", "CLDN5"], "SMC/Fibroblast": ["ACTA2", "MYH11", "TAGLN", "DCN", "LUM", "COL1A1"],
         "Macrophage/Myeloid": ["CD68", "LYZ", "CD14", "AIF1", "ITGAM", "C1QA", "C1QB"], "T/NK": ["CD3D", "CD3E", "TRAC", "CD8A", "NKG7", "GNLY"],
         "B/Plasma": ["CD79A", "MS4A1", "CD19", "IGHG1", "MZB1"], "Mast": ["TPSAB1", "CPA3", "MS4A2"]}


def gene_modules():
    from pyrinplaque import modules as M
    s = M.scoring_sets("config/gene_modules.yaml")
    sets = {"PYRIN_FULL": s["PYRIN_FULL"], "PYRIN_BACKBONE": s["PYRIN_BACKBONE"], "NLRP3_FULL": s["NLRP3_FULL"],
            "NLRP3_BACKBONE": s["NLRP3_BACKBONE"], "EFFECTOR_CYTOKINE_ARM": ["CASP1", "IL1B", "IL18"],
            "EFFECTOR_LYSIS_ARM": ["GSDMD", "GSDME", "NINJ1"], "NONCANONICAL_CASP4_5": ["CASP4", "CASP5"],
            "LEGACY_GENERIC_8GENE": s["GENERIC_PYROPTOSIS"],
            "MYELOID_MARKER_bulk": ["LYZ", "CD68", "CD14", "FCGR3A", "LST1", "C1QA", "C1QB", "C1QC", "MS4A7", "TYROBP"]}
    genes = sorted({g for v in sets.values() for g in v})
    df = pd.DataFrame({k: [int(g in v) for g in genes] for k, v in sets.items()}, index=pd.Index(genes, name="gene"))
    df["note"] = ["added in revision (lysis arm)" if g == "NINJ1" else "bulk symbol DFNA5" if g == "GSDME" else "" for g in genes]
    return df.reset_index()


def tables_spec():
    meta = pd.read_csv("data/metadata/gse159677_sample_metadata.csv")
    wrong = pd.read_csv("data/metadata/gse159677_sample_metadata_ORIGINAL_ERRONEOUS.csv")
    par = yaml.safe_load(open("config/ode_parameters_v2.yaml"))["baseline"]
    sigs = pd.DataFrame([{"signature": k, "genes": ", ".join(v), "level": "myeloid sub-cluster"} for k, v in SIG.items()]
                        + [{"signature": k, "genes": ", ".join(v), "level": "broad compartment"} for k, v in BROAD.items()])
    rd = lambda n, **kw: pd.read_csv(os.path.join(T, n), **kw)
    return [
        ("S1", "Gene-module membership (revised; 1 = member)", [("", gene_modules())]),
        ("S2", "Sample verification: depositor aggregation metadata, barcode matching, depth subsampling and ambient profiles",
         [("Verified sample metadata", meta), ("Original (erroneous) suffix mapping", wrong),
          ("Barcode matching, depth ratio and ambient profile", rd("R02_sample_barcode_verification_and_soup.csv"))]),
        ("S3", "Doublet detection", [("By sample", rd("R03a_doublets_by_sample.csv")), ("By compartment", rd("R03b_doublets_by_compartment.csv")),
                                     ("scDblFinder (rows) vs Scrublet (columns)", rd("R03c_doublet_method_agreement.csv"))]),
        ("S4", "Compartment summary of singlets and module scores", [("Summary incl. delta definitions", rd("R05_compartment_summary_revised.csv")),
                                                                  ("Mean module scores", rd("R11_module_scores_by_compartment.csv"))]),
        ("S5", "Sequencing-depth analyses", [("UMI counts of MEFV+ vs MEFV- cells", rd("R07a_nUMI_MEFVpos_vs_neg.csv")),
                                            ("Complementary log-log models with log(UMI) offset", rd("R07b_cloglog_glm_depth_offset.csv")),
                                            ("Detection within global UMI quintiles", rd("R07c_depth_stratified_detection.csv")),
                                            ("Fixed-depth (2,000 UMI) downsampling", rd("R08_fixed_depth_2000UMI.csv")),
                                            ("Depth-adjusted module scores (OLS)", rd("R08b_depth_adjusted_module_scores.csv"))]),
        ("S6", "Ambient-RNA analyses", [("MEFV detection: observed, corrected and expected from ambient RNA", rd("R09a_ambient_MEFV.csv")),
                                       ("Contamination by sample", rd("R09b_ambient_by_sample.csv")),
                                       ("Module scores after SoupX/DecontX correction", rd("R09c_ambient_corrected_scores.csv"))]),
        ("S7", "Annotation signatures and myeloid sub-cluster signature scores", [("Signatures", sigs),
                                                                                ("Sub-cluster signature z scores", rd("R12a2_subcluster_signature_z.csv"))]),
        ("S8", "Robustness: random gene sets, leave-one-sample-out, patient- and sample-level values",
         [("Expression-matched random gene sets (200)", rd("R10a_random_geneset_null.csv")), ("Leave-one-sample-out", rd("R10b_leave_one_sample_out.csv")),
          ("Patient-level compartment values", rd("R06b_patient_compartment_values.csv")), ("Sample-level compartment values", rd("R06a_sample_compartment_values.csv"))]),
        ("S9", "Myeloid subsets and lineages", [("Sub-clusters", rd("R12a_myeloid_subsets.csv")), ("Lineages (cluster-based and cell-level)", rd("R12b_myeloid_lineages.csv")),
                                               ("Sample x lineage", rd("R12c_sample_by_lineage_MEFV.csv")), ("MEFV co-expression", rd("R12d_MEFV_coexpression_myeloid.csv")),
                                               ("Lineage rate ratios (cloglog)", rd("R12e_lineage_cloglog_glm.csv")),
                                               ("Original exploratory MEFV-high cluster", rd("R12f_v1_MEFVhigh_cluster_mapping.csv"))]),
        ("S10", "Paired core (AC) vs adjacent (PA) comparisons with verified labels", [("Myeloid compartment", rd("R13a_myeloid_AC_vs_PA_corrected.csv")),
                                                                                         ("Lineage-standardized MEFV detection", rd("R13b_region_MEFV_lineage_standardized.csv")),
                                                                                         ("Region models", rd("R13c_region_glm_lineage_adjusted.csv")),
                                                                                         ("All compartments", rd("R13d_all_compartments_AC_vs_PA_corrected.csv"))]),
        ("S11", "Reproduction of the original single-cell pipeline", [("", rd("R01_v1_reproduction_check.csv"))]),
        ("S12", "Bulk gene availability (GSE120521)", [("", rd("R20_bulk_gene_availability.csv"))]),
        ("S13", "Bulk sample-level module scores and residuals", [("", rd("R21_bulk_module_scores.csv"))]),
        ("S14", "Paired bulk results with myeloid residualization", [("", rd("R22_bulk_paired_results.csv"))]),
        ("S15", "Bulk effect sizes versus single-cell myeloid enrichment", [("", rd("R23_bulk_effect_vs_myeloid_enrichment.csv"))]),
        ("S16", "Two-signal model parameters (baseline, dimensionless)", [("", pd.DataFrame({"parameter": list(par), "value": list(par.values())}))]),
        ("S17", "Two-signal model scenario readouts", [("", rd("R30_ode2_scenario_readouts.csv"))]),
        ("S18", "Two-signal model: equal-size perturbations", [("", rd("R31_ode2_matched_perturbations.csv"))]),
        ("S19", "Two-signal model: one-at-a-time sensitivity (+/-20%)", [("", rd("R32_ode2_oat_sensitivity.csv"))]),
        ("S20", "Two-signal model: global sensitivity (LHS, PRCC)", [("", rd("R33_ode2_global_prcc.csv"))]),
        ("S21", "Original model: structural symmetry check", [("", rd("R34_ode1_symmetry_check.csv"))]),
        ("S22", "Priority score components (revised)", [("", rd("R42_priority_components_revised.csv"))]),
        ("S23", "Priority score change log (original -> revised)", [("", rd("R40_priority_change_log.csv"))]),
        ("S24", "Caveat-aware priority table (revised)", [("", rd("R41_priority_table_revised.csv"))]),
    ]


def write_xlsx(spec, path):
    from openpyxl.styles import Font, Alignment
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        idx = pd.DataFrame([{"Table": k, "Title": t, "Blocks": "; ".join(b for b, _ in blocks if b) or "single table"}
                            for k, t, blocks in spec])
        idx.to_excel(xw, sheet_name="Index", index=False)
        for k, title, blocks in spec:
            row = 2
            ws_name = k
            for btitle, df in blocks:
                df = df.copy()
                for c in df.select_dtypes("float").columns:
                    df[c] = df[c].round(6)
                df.to_excel(xw, sheet_name=ws_name, index=False, startrow=row + (1 if btitle else 0))
                ws = xw.sheets[ws_name]
                if btitle:
                    ws.cell(row=row + 1, column=1, value=btitle).font = Font(bold=True, italic=True)
                row += len(df) + (4 if btitle else 3)
            ws = xw.sheets[ws_name]
            ws.cell(row=1, column=1, value=f"Supplementary Table {k}. {title}").font = Font(bold=True, size=12)
            for col in ws.columns:
                width = min(60, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
        ws = xw.sheets["Index"]
        for col, w in zip("ABC", (8, 95, 90)):
            ws.column_dimensions[col].width = w
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True)


FIGS = [
    ("S1", "FigureS1_umap_modules.png", "UMAP of all singlets coloured by broad compartment (left) and by cell-level PYRIN_BACKBONE, NLRP3_BACKBONE and cytokine-arm scores. Module scores are relative transcriptional summaries and do not measure inflammasome activation."),
    ("S2", "FigureS2_gene_dotplot.png", "Detection (dot size) and scaled mean expression (colour) of pyrin-backbone, NLRP3-backbone and shared-effector genes by compartment (singlets). Shared effectors cannot attribute activity to a specific sensor."),
    ("S3", "FigureS3_doublet_ambient.png", "Doublet and ambient-RNA controls. (A) Predicted doublets by sample with scDblFinder (primary) and Scrublet (sensitivity). (B) MEFV detection by compartment: observed in singlets, after SoupX and DecontX correction, and the detection expected from ambient RNA alone (contamination fraction x UMI count x MEFV fraction of the empty-droplet profile); symmetric log axis."),
    ("S4", "FigureS4_depth_controls.png", "Sequencing-depth controls. (A) UMI counts per cell of MEFV-negative and MEFV-positive myeloid cells by sample; horizontal lines indicate medians. (B) MEFV detection by compartment after downsampling every cell with at least 2,000 UMIs to exactly 2,000 UMIs."),
    ("S5", "FigureS5_bulk_gene_heatmap.png", "Gene-level heatmap of the paired bulk samples (GSE120521; row z scores of log2[FPKM+1]), grouped by pyrin backbone, NLRP3 backbone, cytokine arm, lysis arm (GSDME mapped from DFNA5), caspase-4/5, adaptor and myeloid-marker genes."),
    ("S6", "FigureS6_myeloid_signatures.png", "Signature scores (z across sub-clusters) used to label the Harmony-integrated myeloid sub-clusters; each sub-cluster was assigned the signature with the highest standardized score."),
    ("S7", "FigureS7_random_null.png", "Observed compartment PYRIN_BACKBONE means compared with 200 expression-bin-matched random nine-gene sets (mean +/- 1.96 SD) and the empirical percentile of the observed value."),
    ("S8", "FigureS8_v1_symmetry.png", "Superseded model of the original submission: equal shifts of the priming input and of the pyrin activation threshold produce identical outputs because both enter only through their difference; the original scenario contrast therefore reflected perturbation size."),
]

METHODS = """## S-M1. Data acquisition and verification of sample labels

GSE159677: the aggregate filtered feature-barcode matrix (GSE159677_AGGREGATEMAPPED-tisCAR6samples_featurebcmatrixfiltered.tar.gz; SHA-256 8cd0b7f8...), the six per-sample molecule_info.h5 files and the depositors' aggregation metadata (GSM4837528_Aggregated.Sample.Meta.txt.gz) were downloaded from GEO. The metadata assign barcode suffix -1 to Patient 1 PA, -2 to Patient 1 AC, -3 to Patient 2 PA, -4 to Patient 2 AC, -5 to Patient 3 PA and -6 to Patient 3 AC. For each sample, cell barcodes flagged as passing filter in molecule_info.h5 matched 100% of the quality-controlled barcodes of the corresponding suffix and <0.5% of any other suffix (Supplementary Table S2). The original analysis had assigned suffixes in GSM accession order (GSM4837523 to -1, and so on), which interchanged AC and PA within each patient. Per-cell UMI totals in the aggregate were 66-100% of the full-depth totals in molecule_info, consistent with read subsampling to equal mapped depth during aggregation. GSE120521: the FPKM matrix and the series matrix were downloaded; Patient 1 = MB1 (stable)/MB2 (unstable), Patient 2 = MB3/MB4, Patient 3 = MB5/MB6 and Patient 4 = MB9/MB10, as given in the GEO sample characteristics.

## S-M2. Quality control, clustering and compartment annotation

Quality control, normalization, highly variable gene selection, principal component analysis, neighbor graph, Leiden clustering (resolution 0.6) and UMAP followed the original pipeline exactly, which was reproduced before any change (compartment means within 0.001 and identical MEFV detection; Supplementary Table S11). Broad compartments were assigned per Leiden cluster by the highest mean Scanpy score among six marker signatures (Supplementary Table S7).

## S-M3. Doublets and ambient RNA

scDblFinder (R; random artificial doublets, default expected rate of approximately 1% per 1,000 cells) was run per sample on raw counts and defined the analysis set of singlets. Scrublet (Scanpy implementation) was run per sample with an expected doublet rate of 0.8% per 1,000 recovered cells as a sensitivity method. For ambient RNA, each sample's molecule_info.h5 was used to count UMIs for every barcode; non-cell barcodes with 1-100 UMIs defined the empty-droplet (soup) profile, and non-cell barcodes with 10-100 UMIs formed the DecontX background matrix. SoupX was run with the empty-droplet profile and Leiden clusters (autoEstCont), and corrected counts were obtained with adjustCounts (rounded to integers). DecontX (celda) was run with Leiden clusters and the background matrix. The expected number of ambient MEFV UMIs per cell was lambda = rho x UMI count x p_soup(MEFV), and the probability of at least one ambient MEFV UMI was 1 - exp(-lambda).

## S-M4. Module scores, specificity delta and depth models

Module scores used Scanpy score_genes (50 control genes per target gene, 25 expression bins) on log-normalized singlets. Cell-level specificity: z_P = (P - mean(P))/SD(P) and z_N = (N - mean(N))/SD(N) across all singlets, where P and N are the PYRIN_BACKBONE and NLRP3_BACKBONE scores; the delta of a group G is mean over cells in G of (z_P - z_N). Because the z scores have mean zero over all cells, the sum over compartments of n_c x delta_c equals zero. Compartment-mean standardization (reported for comparison) standardizes the six compartment means across compartments (SD with n - 1). Depth: logit-free complementary log-log binomial models, log(-log(1 - p)) = log(UMI) + X beta, imply p = 1 - exp(-UMI x exp(X beta)), so exp(beta) is a ratio of detection rates per UMI. Models included compartment (or lineage, or region) and sample (or patient) fixed effects. Fixed-depth analysis used scanpy.pp.downsample_counts to 2,000 UMIs per cell, followed by normalization, log transformation and re-scoring.

## S-M5. Myeloid sub-clustering

Myeloid singlets were re-processed: 2,000 highly variable genes selected with sample as batch key, scaling, 30 principal components, Harmony integration across the six samples, 15-nearest-neighbor graph, Leiden clustering at resolution 0.6 and UMAP. Each sub-cluster was scored for the signatures in Supplementary Table S7; signature means were standardized across sub-clusters, and each sub-cluster received the label of the highest-scoring signature. Sub-clusters were grouped into monocyte, macrophage, dendritic-cell, neutrophil and mixed lineages. A cluster-independent cell-level lineage call assigned each cell to the lineage with the highest maximum signature score. Monocyte-marker co-expression compared MEFV-positive with MEFV-negative myeloid cells using Haldane-corrected odds ratios and Fisher exact tests.

## S-M6. Bulk analysis

Duplicate symbols were collapsed by the maximum FPKM; DFNA5 was mapped to GSDME. Module scores were means of gene-level z scores of log2(FPKM + 1) across the eight samples. Paired differences used t-distribution 95% confidence intervals (3 degrees of freedom). Residualization: each module score was regressed on the myeloid-marker score across the eight samples, and paired differences of residuals were summarized. Single-cell myeloid enrichment of a module was the mean over its genes of log2[(mean CP10k in myeloid + 0.01)/(mean CP10k in other cells + 0.01)].

## S-M7. Two-signal model

States (population level): inactive pyrin P_i, active pyrin P_a, ASC specks A, active caspase-1 C, GSDMD-N pores G, pro-IL-1beta B_p, intracellular mature IL-1beta B_m, extracellular IL-1beta B_o, the corresponding IL-18 states E_p, E_m and E_o, the lysed fraction L, and two cumulative release counters. With s = 1 - L, sat(x) = x/(1 + x/K_sat), lysis hazard lambda = k_lys x G^h/(K_L^h + G^h) and gate drive = max(0, S2 - theta):

dP_i/dt = s(sigma_0 + sigma_1 x Pi) - k_act x drive x (1 - inh) x P_i + delta x P_a - d_P x P_i - lambda x P_i

dP_a/dt = k_act x drive x (1 - inh) x P_i - delta x P_a - k_ASC x sat(P_a) - lambda x P_a

dA/dt = k_ASC x sat(P_a) - delta x A - lambda x A; dC/dt = k_C1 x sat(A) - delta x C - lambda x C

dG/dt = s x k_G x sat(C) - r_G x G - lambda x G

dB_p/dt = s x sigma_B x Pi - k_mat x sat(C) x B_p - d_pro x B_p - lambda x B_p

dB_m/dt = k_mat x sat(C) x B_p - k_rel x G x B_m - d_m x B_m - lambda x B_m; dB_o/dt = k_rel x G x B_m + lambda x B_m - c x B_o

IL-18 follows the same equations with constitutive synthesis sigma_E in place of sigma_B x Pi; dL/dt = lambda x s.

Initial conditions are the primed steady state before signal 2: P_i(0) = (sigma_0 + sigma_1 x Pi)/d_P, B_p(0) = sigma_B x Pi/d_pro, E_p(0) = sigma_E/d_pro, all other states 0. Parameters are listed in Supplementary Table S16. Partial rank correlation coefficients followed Marino et al.; timing outputs were analyzed only in parameter sets in which the gate opened (signal 2 > threshold; 887 of 1,000 sets).

## S-M8. Priority score

The original candidate inputs were rebuilt exactly from the published component table (the original input file was not in the repository) and the original scores were reproduced. Component values were then changed only where the revised analyses changed the evidence; every change and its reason is listed in Supplementary Table S23. Weights and grade thresholds were not changed.

## S-M9. Software and reproducibility

Python 3.13 with Scanpy 1.12.1, AnnData 0.12.19, NumPy 2.4.6, SciPy 1.18.0, pandas 2.3.3, statsmodels 0.15.0, harmonypy 0.0.10 and matplotlib; R 4.5 with scDblFinder 1.24.10, SoupX 1.6.2 and celda 1.26.0. Scripts in the public repository (scripts/10-16, 05b and 20-24) regenerate every table (results/revision/tables) and figure (results/revision/figures) from the GEO downloads; the manuscript numbers are filled programmatically from these tables (tools/build_manuscript.py)."""

RESULTS = """## S-R1. Sample verification and technical controls

All six samples were verified by barcode matching. scDblFinder removed 10.0% of cells, with higher rates in the larger AC libraries; Scrublet flagged fewer cells, most of which were also scDblFinder doublets (Supplementary Table S3). Ambient contamination was low and the ambient profiles contained very little MEFV, so ambient RNA could explain only a small fraction of MEFV detections (Supplementary Table S6; Supplementary Figure S3).

## S-R2. Depth

MEFV-positive cells had larger libraries than MEFV-negative cells overall, largely because myeloid cells had larger libraries; within myeloid cells the difference was modest and inconsistent across samples. Myeloid enrichment persisted with a log-UMI offset, within UMI quintiles and at fixed depth (Supplementary Table S5; Supplementary Figure S4).

## S-R3. Myeloid subsets and regions

Classical monocytes had the highest MEFV detection in every patient; TREM2+ lipid-associated macrophages were concentrated in AC and rarely expressed MEFV. Lineage standardization narrowed but did not eliminate the lower MEFV detection in AC (Supplementary Tables S9-S10).

## S-R4. Bulk and model

Residualization on the myeloid-marker score attenuated every module by at least 80%, and bulk effect sizes increased with module myeloid specificity (Supplementary Tables S14-S15). In the two-signal model, priming governed IL-1beta magnitude and the IL-1beta:IL-18 ratio, whereas the pyrin gate, caspase-1 and GSDMD steps governed timing (Supplementary Tables S17-S20)."""


def write_docx(spec, path):
    d = DB.new_document(line_numbers=False)
    d.styles["Normal"].paragraph_format.line_spacing = 1.15
    d.styles["Normal"].paragraph_format.space_after = Pt(6)

    def para(text, style="Normal", size=None):
        p = d.add_paragraph(style=style)
        for t, f in DB.parse_inline(text):
            DB._fmt_run(p.add_run(t), f, size)
        return p
    para("**Supplementary Material**")
    para("**Pyrin versus NLRP3: Discriminative Cell-State Mapping of Inflammasome Sensor Programs in Human Atherosclerotic Plaque**")
    para("Contents: Supplementary Methods (S-M1-S-M9); Supplementary Results (S-R1-S-R4); Supplementary Figures S1-S8; "
         "index of Supplementary Tables S1-S24 (provided as Supplementary_Tables.xlsx, one worksheet per table); "
         "Supplementary Appendix: candidate summaries for the priority score.")
    para("Supplementary Methods", "Heading 1")
    for blk in METHODS.split("\n\n"):
        if blk.startswith("## "):
            para(blk[3:], "Heading 2")
        else:
            para(blk)
    para("Supplementary Results", "Heading 1")
    for blk in RESULTS.split("\n\n"):
        if blk.startswith("## "):
            para(blk[3:], "Heading 2")
        else:
            para(blk)
    para("Supplementary Figures", "Heading 1")
    for k, fn, leg in FIGS:
        d.add_picture(os.path.join(FIG, fn), width=Inches(6.5))
        para(f"**Supplementary Figure {k}.** {leg}", size=10)
    para("Supplementary Tables (index)", "Heading 1")
    DB.add_table(d, ["Table", "Title", "Contents"],
                 [[k, t, "; ".join(b for b, _ in blocks if b) or "single table"] for k, t, blocks in spec],
                 [0.6, 3.0, 2.9], note="All tables are provided in Supplementary_Tables.xlsx with full numeric precision.")
    para("Supplementary Appendix. Candidate summaries for the caveat-aware priority score", "Heading 1")
    para("These summaries support research prioritization only; they are not validated target, drug, causal or clinical recommendations.")
    pt = pd.read_csv(os.path.join(T, "R41_priority_table_revised.csv"))
    for i, r in pt.reset_index(drop=True).iterrows():
        para(f"**{i + 1}. {r.candidate_name}** (final score {r.final_priority_score:.2f}; {r.confidence_grade}; raw positive "
             f"{r.raw_positive_score:.2f}, penalties {r.total_penalty:.2f})")
        for lab, fld in (("Supporting evidence", "strongest_supporting_evidence"), ("Counter-evidence", "strongest_counterevidence"),
                         ("Recommended validation", "recommended_next_validation"), ("Appropriate claim", "safe_claim"),
                         ("Claims to avoid", "claims_to_avoid")):
            para(f"*{lab}:* {r[fld]}")
    d.save(path)


def main():
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    spec = tables_spec()
    write_xlsx(spec, os.path.join(out, "Supplementary_Tables_revised.xlsx"))
    write_docx(spec, os.path.join(out, "Supplementary_Material_revised.docx"))
    print("tables:", len(spec))


if __name__ == "__main__":
    main()
