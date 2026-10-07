"""Build the revised manuscript (.docx) from a light markup file, as a clean copy and as a
tracked-changes copy against the originally submitted manuscript.

Markup (one block per paragraph, blank line between blocks):
  # Heading 1 / ## Heading 2 / ### Heading 3
  plain paragraph with *italic*, **bold**, ^superscript^ and ~subscript~
  [[TABLE:key]]   placeholder replaced by a table supplied in python
  [[PAGEBREAK]]
"""
import copy
import datetime
import difflib
import re

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

AUTHOR = "İbrahim Halil Tanboğa"
DATE = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).strftime("%Y-%m-%dT%H:%M:%SZ")
_ID = [1000]
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\^.+?\^|~.+?~)")


def _next_id():
    _ID[0] += 1
    return str(_ID[0])


# ---------------------------------------------------------------- markup -> runs
def parse_inline(text):
    """Return list of (text, fmt) with fmt in {'', 'b', 'i', 'sup', 'sub'}."""
    out = []
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], "b"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], "i"))
        elif part.startswith("^") and part.endswith("^"):
            out.append((part[1:-1], "sup"))
        elif part.startswith("~") and part.endswith("~"):
            out.append((part[1:-1], "sub"))
        else:
            out.append((part, ""))
    return out


def parse_markup(src):
    blocks = []
    for raw in re.split(r"\n\s*\n", src.strip()):
        b = raw.strip().replace("\n", " ")
        if not b:
            continue
        if b.startswith("[[TABLE:"):
            blocks.append({"kind": "table", "key": b[8:-2]})
        elif b == "[[PAGEBREAK]]":
            blocks.append({"kind": "pagebreak"})
        else:
            m = re.match(r"^(#{1,3})\s+(.*)$", b)
            if m:
                blocks.append({"kind": f"h{len(m.group(1))}", "runs": parse_inline(m.group(2))})
            else:
                blocks.append({"kind": "p", "runs": parse_inline(b)})
    return blocks


# ---------------------------------------------------------------- document setup
def new_document(line_numbers=True):
    d = docx.Document()
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.comments = ""
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    for lvl, size in ((1, 13), (2, 12), (3, 12)):
        h = d.styles[f"Heading {lvl}"]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = lvl == 3
        h.font.color.rgb = RGBColor(0, 0, 0)
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), "Times New Roman")
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if line_numbers:
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        # schema order: ... pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols, ..., docGrid
        anchor = None
        for tag in ("w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                    "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid"):
            anchor = sec._sectPr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(ln)
        else:
            sec._sectPr.append(ln)
    zoom = d.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    # page number footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE")
    return d


def _field(paragraph, instr):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._r.append(b)
    r._r.append(t)
    r._r.append(e)


def _fmt_run(run, fmt, size=None):
    if fmt == "b":
        run.bold = True
    elif fmt == "i":
        run.italic = True
    elif fmt == "sup":
        run.font.superscript = True
    elif fmt == "sub":
        run.font.subscript = True
    if size:
        run.font.size = Pt(size)


def add_block(d, block, tables=None, size=None):
    k = block["kind"]
    if k == "pagebreak":
        d.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        return
    if k == "table":
        tables[block["key"]](d)
        return
    p = d.add_paragraph(style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}.get(k, "Normal"))
    for text, fmt in block["runs"]:
        _fmt_run(p.add_run(text), fmt, size)
    return p


# ---------------------------------------------------------------- tables
def add_table(d, header, rows, widths_in, caption=None, note=None, font=9):
    if caption:
        p = d.add_paragraph()
        for text, fmt in parse_inline(caption):
            _fmt_run(p.add_run(text), fmt)
        p.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Inches(widths_in[i])
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
        for text, fmt in parse_inline(h):
            r = c.paragraphs[0].add_run(text)
            r.bold = True
            _fmt_run(r, fmt, font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Inches(widths_in[i])
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
            for text, fmt in parse_inline(str(v)):
                _fmt_run(cells[i].paragraphs[0].add_run(text), fmt, font)
    if note:
        p = d.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        for text, fmt in parse_inline(note):
            _fmt_run(p.add_run(text), fmt, 9)
    d.add_paragraph()
    return t


# ---------------------------------------------------------------- tracked changes
def _tok(runs):
    """Split (text, fmt) runs into word tokens that keep their trailing whitespace."""
    out = []
    for text, fmt in runs:
        for m in re.finditer(r"\S+\s*|\s+", text):
            out.append((m.group(0), fmt))
    return out


def _run_xml(text, fmt, deleted=False):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if fmt == "b":
        rpr.append(OxmlElement("w:b"))
    if fmt == "i":
        rpr.append(OxmlElement("w:i"))
    if fmt in ("sup", "sub"):
        v = OxmlElement("w:vertAlign")
        v.set(qn("w:val"), "superscript" if fmt == "sup" else "subscript")
        rpr.append(v)
    if len(rpr):
        r.append(rpr)
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _wrap(kind, runs_xml):
    w = OxmlElement(f"w:{kind}")
    w.set(qn("w:id"), _next_id())
    w.set(qn("w:author"), AUTHOR)
    w.set(qn("w:date"), DATE)
    for r in runs_xml:
        w.append(r)
    return w


def _mark_paragraph(p, kind):
    pPr = p._p.get_or_add_pPr()
    rpr = pPr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        pPr.append(rpr)
    m = OxmlElement(f"w:{kind}")
    m.set(qn("w:id"), _next_id())
    m.set(qn("w:author"), AUTHOR)
    m.set(qn("w:date"), DATE)
    rpr.insert(0, m)


def _emit(p, toks, kind=None):
    if not toks:
        return
    # merge consecutive tokens with same fmt
    merged = []
    for text, fmt in toks:
        if merged and merged[-1][1] == fmt:
            merged[-1] = (merged[-1][0] + text, fmt)
        else:
            merged.append((text, fmt))
    xs = [_run_xml(t, f, deleted=(kind == "del")) for t, f in merged]
    if kind:
        p._p.append(_wrap(kind, xs))
    else:
        for x in xs:
            p._p.append(x)


def tracked_paragraph(d, style, old_runs, new_runs):
    p = d.add_paragraph(style=style)
    a, b = _tok(old_runs), _tok(new_runs)
    sm = difflib.SequenceMatcher(None, [x[0].strip() for x in a], [x[0].strip() for x in b], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            _emit(p, b[j1:j2])
        else:
            _emit(p, a[i1:i2], "del")
            _emit(p, b[j1:j2], "ins")
    return p


def deleted_paragraph(d, style, runs):
    p = d.add_paragraph(style=style)
    _emit(p, _tok(runs), "del")
    _mark_paragraph(p, "del")
    return p


def inserted_paragraph(d, style, runs):
    p = d.add_paragraph(style=style)
    _emit(p, _tok(runs), "ins")
    _mark_paragraph(p, "ins")
    return p


def _plain(runs):
    return "".join(t for t, _ in runs)


def align_blocks(old, new, threshold=0.35):
    """Needleman-Wunsch alignment of paragraph lists by token similarity.
    Returns list of (old_index|None, new_index|None)."""
    n, m = len(old), len(new)
    ow = [_plain(b["runs"]).split() for b in old]
    nw = [_plain(b["runs"]).split() for b in new]

    def sim(i, j):
        if not ow[i] or not nw[j]:
            return 0.0
        return difflib.SequenceMatcher(None, ow[i], nw[j], autojunk=False).ratio()

    S = [[0.0] * (m + 1) for _ in range(n + 1)]
    B = [[None] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        B[i][0] = "u"
    for j in range(1, m + 1):
        B[0][j] = "l"
    cache = {}
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            s = cache.setdefault((i, j), sim(i - 1, j - 1))
            diag = S[i - 1][j - 1] + (s if s >= threshold else -1e9)
            up, left = S[i - 1][j], S[i][j - 1]
            best = max(diag, up, left)
            S[i][j] = best
            B[i][j] = "d" if best == diag else ("u" if best == up else "l")
    i, j, path = n, m, []
    while i > 0 or j > 0:
        c = B[i][j]
        if c == "d":
            path.append((i - 1, j - 1))
            i, j = i - 1, j - 1
        elif c == "u":
            path.append((i - 1, None))
            i -= 1
        else:
            path.append((None, j - 1))
            j -= 1
    return path[::-1]


def blocks_from_docx(path, stop_at=None):
    """Original manuscript paragraphs as markup-like blocks (style -> kind, italics kept)."""
    src = docx.Document(path)
    blocks = []
    for p in src.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if stop_at and text.startswith(stop_at):
            break
        sty = p.style.name
        kind = "h1" if sty == "Heading 1" else "h2" if sty == "Heading 2" else "h3" if sty == "Heading 3" else "p"
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            fmt = "i" if r.italic else "b" if r.bold else "sup" if r.font.superscript else "sub" if r.font.subscript else ""
            runs.append((r.text, fmt))
        blocks.append({"kind": kind, "runs": runs})
    return blocks
